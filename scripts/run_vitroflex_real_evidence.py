"""
Ejecución real de evidencia Vitroflex: plantillas copiadas, PDFs, validación texto/render,
API Flask y limpieza de temporales (LibreOffice interno).

Salida: tests/evidence_vitroflex/real_run/
"""

from __future__ import annotations

import io
import json
import re
import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "tests" / "evidence_vitroflex" / "real_run"
PDF_DIR = OUT / "pdfs"
PNG_DIR = OUT / "renders"
CMP_DIR = OUT / "compare_template"


def _ensure_dirs() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    PDF_DIR.mkdir(parents=True, exist_ok=True)
    PNG_DIR.mkdir(parents=True, exist_ok=True)
    CMP_DIR.mkdir(parents=True, exist_ok=True)


def _write_excel_two_cols(path: Path, n: int = 8) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["NOMBRE TRABAJADOR", "NO. IMSS"])
    for i in range(n):
        ws.append([f"Persona Ejemplo {i+1}", f"{11000000000 + i:011d}"])
    wb.save(path)


def _write_excel_four_cols(path: Path, n: int = 8) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["NOMBRE TRABAJADOR", "NO. IMSS", "ACTIVIDAD A REALIZAR", "TEL. EMERGENCIA"])
    for i in range(n):
        ws.append(
            [
                f"Persona Full {i+1}",
                f"{12000000000 + i:011d}",
                "Aux. de limpieza",
                "81 2183 9413",
            ]
        )
    wb.save(path)


def _pdf_validate(path: Path, *, kind: str) -> tuple[str, list[str]]:
    """Retorna (estado, notas)."""
    import fitz

    notes: list[str] = []
    doc = fitz.open(str(path))
    try:
        if len(doc) == 0:
            return "FAIL", ["sin páginas"]
        text_all = ""
        for i in range(len(doc)):
            text_all += doc.load_page(i).get_text() or ""
        if len(text_all.strip()) < 80:
            return "FAIL", [f"texto muy corto ({len(text_all)} chars)"]
        if "{{FECHA}}" in text_all or "{{PLANTA}}" in text_all or "{{PERMISO" in text_all:
            return "FAIL", ["placeholders sin reemplazar en texto extraído"]
        notes.append(f"páginas={len(doc)}, chars≈{len(text_all)}")
        if kind == "memo":
            for needle in ("Para", "De", "Relaciones", "ATENT"):
                if needle.lower() not in text_all.lower():
                    notes.append(f"aviso: no se encontró '{needle}' en texto")
        if kind == "cr":
            if "{{" in text_all:
                return "FAIL", ["sigue habiendo {{ en PDF"]
        return "OK", notes
    finally:
        doc.close()


def _render_png(pdf_path: Path, stem: str, max_pages: int = 3) -> list[Path]:
    import fitz

    out: list[Path] = []
    doc = fitz.open(str(pdf_path))
    try:
        for i in range(min(max_pages, len(doc))):
            pix = doc.load_page(i).get_pixmap(matrix=fitz.Matrix(2.0, 2.0), alpha=False)
            p = PNG_DIR / f"{stem}_p{i+1}.png"
            pix.save(str(p))
            out.append(p)
    finally:
        doc.close()
    return out


def _gen_pdf_from_bytes(docx_bytes: bytes, name: str) -> Path:
    from modules.vitroflex_docs.libreoffice_pdf import docx_bytes_to_pdf_bytes

    pdf_bytes = docx_bytes_to_pdf_bytes(docx_bytes, suffix=name[:40])
    p = PDF_DIR / f"{name}.pdf"
    p.write_bytes(pdf_bytes)
    return p


def _template_reference_pdf_memo() -> Path | None:
    from modules.vitroflex_docs.libreoffice_pdf import docx_to_pdf
    from modules.vitroflex_docs.template_paths import MEMO_DOCX

    if not MEMO_DOCX.is_file():
        return None
    outp = CMP_DIR / "A_MEMO_plantilla_sin_llenar.pdf"
    docx_to_pdf(MEMO_DOCX, outp)
    return outp


def main() -> int:
    sys.path.insert(0, str(ROOT))
    import os

    os.chdir(ROOT)
    _ensure_dirs()

    from modules.vitroflex_docs.build_document import build_cr_docx_bytes, build_memo_docx_bytes
    from modules.vitroflex_docs.excel_import import DEFAULT_ACTIVIDAD, DEFAULT_TEL, parse_excel_bytes
    from modules.vitroflex_docs.libreoffice_pdf import docx_bytes_to_pdf_bytes, resolve_soffice_path
    from modules.vitroflex_docs.template_paths import CR_DOCX, MEMO_DOCX

    report: list[str] = []
    report.append(f"# Evidencia real Vitroflex — {datetime.now().isoformat(timespec='seconds')}")
    report.append("")
    report.append("## Entorno")
    report.append(f"- LibreOffice: `{resolve_soffice_path()}`")
    report.append(f"- MEMO template: `{MEMO_DOCX}` exists={MEMO_DOCX.is_file()}")
    report.append(f"- CR template: `{CR_DOCX}` exists={CR_DOCX.is_file()}")
    report.append("")

    if not MEMO_DOCX.is_file() or not CR_DOCX.is_file():
        report.append("FAIL: faltan plantillas en vitroflex_templates/")
        (OUT / "REPORTE_EJECUCION.md").write_text("\n".join(report), encoding="utf-8")
        print("Faltan plantillas")
        return 2

    # Excel temporal (equivalente ejemplo 1 / 2)
    x1 = OUT / "_tmp_ejemplo1.xlsx"
    x2 = OUT / "_tmp_ejemplo2.xlsx"
    _write_excel_two_cols(x1, n=10)
    _write_excel_four_cols(x2, n=10)

    rows_e1, nd1, err1 = parse_excel_bytes(x1.read_bytes())
    rows_e2, nd2, err2 = parse_excel_bytes(x2.read_bytes())
    report.append("## Import Excel")
    report.append(f"- Ejemplo 1 (2 cols): filas={len(rows_e1)} needs_defaults={nd1} err={err1}")
    report.append(f"- Ejemplo 2 (4 cols): filas={len(rows_e2)} needs_defaults={nd2} err={err2}")
    for r in rows_e1:
        if not r.get("actividad"):
            r["actividad"] = DEFAULT_ACTIVIDAD
        if not r.get("tel"):
            r["tel"] = DEFAULT_TEL
    report.append("")

    workers3 = [
        {"nombre": "Ana López Martínez", "imss": "12345678901", "actividad": "Aux.", "tel": "81 2183 9413"},
        {"nombre": "Luis Hernández", "imss": "10987654321", "actividad": "Aux.", "tel": "81 2183 9413"},
        {"nombre": "Rosa Fuentes", "imss": "55555555555", "actividad": "Apoyo", "tel": "81 2000 0000"},
    ]
    big = [
        {
            "nombre": f"Trabajador {i+1}",
            "imss": f"{10000000000 + i:011d}",
            "actividad": "Aux. de limpieza",
            "tel": "81 2183 9413",
        }
        for i in range(29)
    ]

    cases: list[tuple[str, str, dict]] = [
        (
            "01_memo_manual_3",
            "memo",
            {
                "fecha": "Garcia, N. L. a 15 de marzo de 2026 — evidencia E2E",
                "p1": "2026-03-10",
                "p2": "2026-04-27",
                "w": workers3,
            },
        ),
        (
            "02_memo_import_ejemplo1",
            "memo",
            {
                "fecha": "Importación tipo ejemplo 1 (2 columnas + defaults)",
                "p1": "2026-06-15",
                "p2": "2026-07-02",
                "w": rows_e1,
            },
        ),
        (
            "03_memo_import_ejemplo2",
            "memo",
            {
                "fecha": "Importación tipo ejemplo 2 (4 columnas)",
                "p1": "2026-01-01",
                "p2": "2026-02-18",
                "w": rows_e2,
            },
        ),
        (
            "04_cr_vitroflex",
            "cr",
            {"fecha": "Fecha CR Vitroflex evidencia", "planta": "Vitroflex", "w": workers3[:2]},
        ),
        (
            "05_cr_mercado_moderno",
            "cr",
            {"fecha": "Fecha CR Mercado Moderno", "planta": "Mercado Moderno", "w": workers3[:2]},
        ),
        (
            "06_caso_una_hoja",
            "memo",
            {
                "fecha": "Caso pocas filas",
                "p1": "2026-03-01",
                "p2": "2026-04-18",
                "w": workers3,
            },
        ),
        (
            "07_caso_multiples_hojas",
            "memo",
            {
                "fecha": "Caso multipágina",
                "p1": "2026-01-01",
                "p2": "2026-02-18",
                "w": big,
            },
        ),
    ]

    results: list[tuple[str, str, dict]] = []

    for stem, k, data in cases:
        try:
            if k == "memo":
                b = build_memo_docx_bytes(
                    fecha_texto=data["fecha"],
                    permiso1_iso=data.get("p1"),
                    permiso2_iso=data.get("p2"),
                    workers=data["w"],
                    template_path=MEMO_DOCX,
                )
            else:
                b = build_cr_docx_bytes(
                    fecha_texto=data["fecha"],
                    planta=data["planta"],
                    workers=data["w"],
                    template_path=CR_DOCX,
                )
            pdf_path = _gen_pdf_from_bytes(b, stem)
            status, notes = _pdf_validate(pdf_path, kind=k)
            pngs = _render_png(pdf_path, stem, max_pages=3 if "multiples" in stem else 2)
            results.append(
                (
                    stem,
                    status,
                    {
                        "pdf": str(pdf_path),
                        "notes": notes,
                        "pngs": [str(x) for x in pngs],
                    },
                )
            )
        except Exception as exc:
            results.append((stem, "FAIL", {"error": str(exc)}))

    report.append("## Resultados por caso")
    for stem, status, info in results:
        report.append(f"### {stem}: **{status}**")
        report.append(f"```json\n{json.dumps(info, ensure_ascii=False, indent=2)}\n```")
        report.append("")

    # Comparación plantilla MEMO sin llenar
    ref = _template_reference_pdf_memo()
    if ref:
        report.append("## Comparación plantilla MEMO (sin llenar)")
        report.append(f"- PDF referencia: `{ref}`")
        _render_png(ref, "A_plantilla_memo_ref", max_pages=1)
        report.append("")

    # API Flask: status + generación
    report.append("## API backend (Flask test_client)")
    try:
        from app import app

        app.config["TESTING"] = True
        c = app.test_client()
        st = c.get("/vitroflex/api/status")
        report.append(f"- GET /vitroflex/api/status → {st.status_code}")
        if st.is_json:
            report.append(f"```json\n{json.dumps(st.get_json(), ensure_ascii=False, indent=2)}\n```")
        import sqlite3
        from app import DB_PATH

        conn = sqlite3.connect(DB_PATH)
        uid = conn.execute("SELECT id FROM users LIMIT 1").fetchone()[0]
        conn.close()
        with c.session_transaction() as sess:
            sess["user_id"] = uid
        st2 = c.get("/vitroflex/api/status")
        report.append(f"- Con sesión: {st2.status_code}")
        if st2.is_json:
            report.append(f"```json\n{json.dumps(st2.get_json(), ensure_ascii=False, indent=2)}\n```")
        payload = {
            "kind": "memo",
            "fecha_texto": "API test",
            "permiso1": "2026-03-10",
            "permiso2": "2026-04-27",
            "workers": workers3,
            "filename": "api_test.pdf",
            "disposition": "attachment",
        }
        r = c.post("/vitroflex/api/generate-pdf", json=payload)
        report.append(f"- POST generate-pdf: {r.status_code}, bytes={len(r.data)}")
        if r.status_code == 200 and r.data[:4] == b"%PDF":
            (OUT / "api_smoke_generate.pdf").write_bytes(r.data)
            report.append("- PDF guardado en `real_run/api_smoke_generate.pdf`")
        else:
            report.append(f"- Error: {r.get_data(as_text=True)[:500]}")
    except Exception as exc:
        report.append(f"- **FAIL API** {exc}")

    report.append("")
    report.append("## Temporales LibreOffice")
    report.append("Los temporales usados por `docx_bytes_to_pdf_bytes` están en `tempfile.TemporaryDirectory` y se eliminan al salir del contexto (ver `libreoffice_pdf.py`).")

    (OUT / "REPORTE_EJECUCION.md").write_text("\n".join(report), encoding="utf-8")
    print("Escrito:", OUT / "REPORTE_EJECUCION.md")
    fails = sum(1 for _, s, _ in results if s != "OK")
    return 1 if fails else 0


if __name__ == "__main__":
    raise SystemExit(main())
